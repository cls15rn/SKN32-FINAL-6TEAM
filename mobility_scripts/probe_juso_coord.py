# mobility_scripts/probe_juso_coord.py — 도로명주소 → 좌표 파이프라인 확인
#
# 무엇을 확인하나
#   ① 주소검색 API → 좌표제공 API 2단계가 실제로 도는가
#   ② 좌표제공 API 가 주는 좌표계가 무엇인가 (문서에 명시가 없다. UTM-K/EPSG:5179 로 알려져 있으나 확인 필요)
#   ③ 변환한 좌표가 맞는가 — **정답지가 우리에게 있다**
#      국가철도공단 표준데이터 XLSX 한 파일에 `역사도로명주소` 와 `역위도`·`역경도` 가 같이 들어 있다.
#      같은 역의 주소를 juso 로 돌려 나온 좌표를 그 위경도와 비교하면 파이프라인 전체가 한 번에 검증된다.
#
# 준비
#   1) business.juso.go.kr 에서 승인키 발급 — ★검색용과 좌표제공용이 **따로**다
#   2) .env 에 JUSO_SEARCH_KEY=... / JUSO_COORD_KEY=...  (같은 키면 둘 다 같은 값으로)
#   3) pip install pyproj   (없으면 좌표계 판별까지만 하고 변환은 건너뛴다)
#
# 실행:  python mobility_scripts/probe_juso_coord.py [--n 8]
import argparse, json, math, os, sys
from pathlib import Path

import requests

REPO = Path(__file__).resolve().parents[1]
SEARCH = "https://business.juso.go.kr/addrlink/addrLinkApi.do"
COORD = "https://business.juso.go.kr/addrlink/addrCoordApi.do"


def env(name, *alts):
    p = REPO / ".env"
    if p.exists():
        for line in p.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            for k in (name, *alts):
                if line.startswith(k + "="):
                    return line.split("=", 1)[1].strip().strip('"')
    return os.environ.get(name)


def data_dir():
    root = env("DATA_DIR")
    return (Path(root) if root else REPO / "data") / "travel"


def haversine(lat1, lng1, lat2, lng2):
    R = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lng2 - lng1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(a))


ap = argparse.ArgumentParser()
ap.add_argument("--n", type=int, default=8, help="표본 역 수")
a = ap.parse_args()

SK, CK = env("JUSO_SEARCH_KEY", "JUSO_KEY"), env("JUSO_COORD_KEY", "JUSO_KEY")
if not SK or not CK:
    sys.exit(".env 에 JUSO_SEARCH_KEY / JUSO_COORD_KEY 가 없다 (business.juso.go.kr 에서 각각 발급)")

# ── 정답지: 국가철도공단 표준데이터 ──
xl = sorted((data_dir() / "raw" / "mobility").glob("*도시철도역사정보*.xlsx"))
if not xl:
    sys.exit("표준데이터 XLSX 가 없다: raw/mobility/*도시철도역사정보*.xlsx")
import openpyxl
ws = openpyxl.load_workbook(xl[-1], read_only=True)[openpyxl.load_workbook(xl[-1], read_only=True).sheetnames[0]]
it = ws.iter_rows(values_only=True)
head = list(next(it))
c = {k: head.index(k) for k in ("역사명", "역위도", "역경도", "역사도로명주소", "운영기관명")}
truth = []
for r in it:
    addr, lat, lng = r[c["역사도로명주소"]], r[c["역위도"]], r[c["역경도"]]
    if not addr or not lat or not str(addr).startswith("서울특별시"):
        continue
    truth.append({"역": r[c["역사명"]], "주소": str(addr).strip(),
                  "lat": float(lat), "lng": float(lng)})
    if len(truth) >= a.n:
        break
print(f"정답지 {xl[-1].name} · 서울 표본 {len(truth)}역\n")

sess = requests.Session()
rows = []
for t in truth:
    r = sess.get(SEARCH, params={"confmKey": SK, "currentPage": 1, "countPerPage": 1,
                                 "keyword": t["주소"], "resultType": "json"}, timeout=20)
    try:
        body = r.json()["results"]
    except Exception:
        print(f"  {t['역']:<10} 검색 응답 파싱 실패: {r.text[:120]}"); continue
    if body["common"]["errorCode"] != "0":
        print(f"  {t['역']:<10} 검색 오류 {body['common']['errorCode']} {body['common']['errorMessage']}"); continue
    if not body["juso"]:
        print(f"  {t['역']:<10} 검색 결과 없음 — {t['주소']}"); continue
    j = body["juso"][0]
    r2 = sess.get(COORD, params={"confmKey": CK, "admCd": j["admCd"], "rnMgtSn": j["rnMgtSn"],
                                 "udrtYn": j["udrtYn"], "buldMnnm": j["buldMnnm"],
                                 "buldSlno": j["buldSlno"], "resultType": "json"}, timeout=20)
    try:
        b2 = r2.json()["results"]
    except Exception:
        print(f"  {t['역']:<10} 좌표 응답 파싱 실패: {r2.text[:120]}"); continue
    if b2["common"]["errorCode"] != "0":
        print(f"  {t['역']:<10} 좌표 오류 {b2['common']['errorCode']} {b2['common']['errorMessage']}"); continue
    if not b2["juso"]:
        print(f"  {t['역']:<10} 좌표 결과 없음"); continue
    k = b2["juso"][0]
    rows.append({**t, "entX": float(k["entX"]), "entY": float(k["entY"]),
                 "juso_addr": j.get("roadAddr")})
    print(f"  {t['역']:<10} entX={k['entX']} entY={k['entY']}")

if not rows:
    sys.exit("\n좌표를 하나도 못 받았다. 승인키와 API 종류(검색/좌표제공)를 확인한다.")

# ── ② 좌표계 판별 ──
x = rows[0]["entX"]; y = rows[0]["entY"]
print("\n[좌표계 판별]")
if 120 < x < 132 and 33 < y < 39:
    crs = "EPSG:4326"; print("  entX/entY 가 경도·위도 범위다 → 이미 WGS84 위경도. 변환 불필요")
elif 1e5 < x < 1.5e6 and 1e5 < y < 2.5e6:
    crs = "EPSG:5179"; print(f"  entX={x:,.0f} entY={y:,.0f} → 평면좌표. GRS80 UTM-K(EPSG:5179)로 본다")
else:
    crs = None; print(f"  판별 불가: entX={x} entY={y} — 값 범위를 보고 좌표계를 직접 정해야 한다")

# ── ③ 정답지 대조 ──
print("\n[정답지 대조] 국가철도공단 역위도·역경도와의 거리")
conv = None
if crs == "EPSG:4326":
    conv = lambda X, Y: (Y, X)
elif crs == "EPSG:5179":
    try:
        from pyproj import Transformer
        tf = Transformer.from_crs("EPSG:5179", "EPSG:4326", always_xy=True)
        conv = lambda X, Y: tuple(reversed(tf.transform(X, Y)))   # (lat, lng)
    except ImportError:
        print("  pyproj 가 없다 → pip install pyproj 후 다시 실행하면 변환·대조까지 된다")

if conv:
    ds = []
    print(f"  {'역':<10}{'거리(m)':>9}   변환 좌표 / 표준데이터")
    for r in rows:
        lat, lng = conv(r["entX"], r["entY"])
        d = haversine(lat, lng, r["lat"], r["lng"])
        ds.append(d)
        print(f"  {r['역']:<10}{d:>9.0f}   {lat:.6f},{lng:.6f} / {r['lat']:.6f},{r['lng']:.6f}")
    ds.sort()
    med = ds[len(ds) // 2]
    print(f"\n  거리 중앙 {med:.0f}m · 최대 {ds[-1]:.0f}m")
    if med < 200:
        print("  → 좌표계와 변환이 맞다. 이 파이프라인으로 카탈로그 좌표를 채울 수 있다.")
        print("     (수십~백여 m 차이는 정상 — juso 는 건물 출입구, 표준데이터는 역사 대표점이다)")
    else:
        print("  → 중앙값이 크다. 좌표계 가정이 틀렸거나 주소 매칭이 어긋났다. entX/entY 원값을 다시 본다.")

out = data_dir() / "raw" / "probe" / "juso_coord_probe.json"
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps({"crs_guess": crs, "rows": rows}, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"\n원본 → {out}")
