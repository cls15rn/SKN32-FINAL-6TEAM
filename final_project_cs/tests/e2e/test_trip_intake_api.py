# -*- coding: utf-8 -*-
"""계획 읽기 — 웹 입구 `/v1/web/trip-intakes`. `[2026-09-27]` 설계서 1주차

★완료 기준(설계서 §8 1주차): 실측 사진이 **6/6 항목, 값 변조 0** 으로 `intake_claims` 에 들어간다.
  여기서는 받아쓰기 모델을 흉내 낸다(사진의 원문을 그대로 돌려준다) — 실제 gemma4:12b 로는 같은 사진에서
  6/6·변조 0·누락 0 을 쟀다(2026-09-27, 1회, 전체+반쪽 둘 56.6초).
★값은 모두 원문의 (줄, 시작, 끝) 근거를 들고 다닌다 — 근거 없는 값 0.

재현:

    python -m pytest tests/e2e/test_trip_intake_api.py -v
"""
from __future__ import annotations

import io
from pathlib import Path
from uuid import UUID

import pytest
from fastapi.testclient import TestClient

from app.infrastructure.db.session import get_connection
from app.presentation.api.app import create_app
from app.modules.travel_ops.trip_api import build_trip_router

from .test_trip_api import api  # noqa: F401 — 픽스처를 그대로 쓴다

PHOTO = Path(__file__).resolve().parents[1] / "fixtures" / "intake" / "seoul_2n3d_plan.png"
#: 사진에 그려 넣은 원문 — 가짜 받아쓰기가 그대로 돌려준다
PHOTO_TEXT = """서울 2박3일 가족여행 (4명)

1일차 · 10월 15일 (수)
09:00 경복궁 관람
12:30 토속촌삼계탕 점심
19:00 명동난타극장 · 예약번호 KY-20931

2일차 · 10월 16일 (목)
10:00 북촌한옥마을 산책
13:00 광장시장 빈대떡
15:30 N서울타워"""
WANT = [("09:00", "경복궁 관람"), ("12:30", "토속촌삼계탕"), ("19:00", "명동난타극장"),
        ("10:00", "북촌한옥마을 산책"), ("13:00", "광장시장 빈대떡"), ("15:30", "N서울타워")]


class Vision:
    """받아쓰기 흉내 — 전체든 반쪽이든 원문을 돌려준다(반쪽 누락 검사가 「빠진 줄 없음」이 되게)."""

    def __init__(self):
        self.calls = 0

    def see(self, prompt, image):
        self.calls += 1
        assert "그대로" in prompt                  # ★구조화를 시키지 않는다 — 받아 적기만
        return PHOTO_TEXT


def _client(vision=None):
    return TestClient(create_app(classifier=lambda _m: {"intent": "other", "issue_code": "other",
                                                          "sentiment": "neutral"},
                                 domain_routers=[build_trip_router(chat_factory=(lambda: vision)
                                                                   if vision else None)]))


def _key(client) -> dict:
    return {"X-User-Key": client.post("/v1/web/session").json()["user_key"]}


def _send(client, headers, *, text="", files=()):
    response = client.post("/v1/web/trip-intakes", headers=headers, data={"text": text},
                           files=[("files", (name, data, "application/octet-stream")) for name, data in files])
    assert response.status_code == 202, response.text
    return client.get(f"/v1/web/trip-intakes/{response.json()['intake_id']}", headers=headers).json()


def _items(view, source=0):
    return [(i["fields"]["starts_at"]["value"] if "starts_at" in i["fields"] else None,
             i["fields"]["title"]["value"] if "title" in i["fields"] else None)
            for i in view["sources"][source]["items"]]


def test_a_photo_plan_is_read_six_of_six_with_every_value_cut_from_the_transcript(api):
    vision = Vision()
    client = _client(vision)
    headers = _key(client)
    view = _send(client, headers, files=[("plan.png", PHOTO.read_bytes())])
    assert view["status"] == "review", view
    assert vision.calls == 3                                   # 전체 + 위·아래 반쪽(누락 검사)
    source = view["sources"][0]
    assert source["kind"] == "image" and source["transcribed"] and source["missing"] == []
    assert _items(view) == WANT                                 # ★6/6
    lines = PHOTO_TEXT.splitlines()
    for item in source["items"]:                                # ★값 변조 0 — 이름은 원문 조각 그대로
        title = item["fields"]["title"]
        ev = title["evidence"]
        assert lines[ev["line"] - 1][ev["start"]:ev["end"]] == title["value"] == ev["text"]
        assert title["method"] == "rule"
    booking = [i["fields"]["booking_no"]["value"] for i in source["items"] if "booking_no" in i["fields"]]
    assert booking == ["KY-20931"]
    assert source["trip"]["party_size"]["value"] == 4 and source["trip"]["nights"]["value"] == 2
    # ★2주차 — 해 없는 「10월 15일」은 오늘(서울) 기준 다가오는 날로 해를 채우고, 확인을 받는다
    assert [(i["day"], i["date"][4:]) for i in source["items"]][:1] == [(1, "-10-15")]
    assert [(i["day"], i["date"][4:]) for i in source["items"]][3:4] == [(2, "-10-16")]
    first_date = source["items"][0]["fields"]["date"]
    assert first_date["evidence"]["how"] == "year_filled" and first_date["needs_review"]


def test_a_chat_style_plan_is_read_into_items(api):
    client = _client()
    view = _send(client, _key(client),
                 text="첫날은 9시에 경복궁 가고, 12시 반에 토속촌에서 점심 먹고 오후 3시에 인사동 가요\n"
                      "둘째 날 10시~12시 한강 카약, 저녁 7시에 명동")
    assert view["status"] == "review" and view["sources"][0]["kind"] == "chat"
    assert _items(view) == [("09:00", "경복궁"), ("12:30", "토속촌"), ("15:00", "인사동"),
                            ("10:00", "한강 카약"), ("19:00", "명동")]
    ends = [i["fields"].get("ends_at", {}).get("value") for i in view["sources"][0]["items"]]
    assert ends[3] == "12:00"


def test_an_unclear_hour_is_marked_for_review_and_unread_lines_are_kept(api):
    client = _client()
    view = _send(client, _key(client), text="3시에 남산타워\n저녁은 아무 데나 괜찮아요")
    fields = [r["field"] for r in view["needs_review"]]
    assert [f for f in fields if f.endswith("starts_at")] == ["items[0].starts_at"]
    # ★2주차 — 날짜가 어디에도 없으면 지어내지 않고 첫날 하나만 묻는다. 장소 조회기가 없으면 「못 정함」으로 남는다
    assert "trip.ask_first_day" in fields and "items[0].date" in fields and "items[0].place" in fields
    lines = view["sources"][0]["lines"]
    assert [line["read"] for line in lines] == [True, False]     # 남은 줄 — 2주차에 LLM 이 위치만 가리킨다


def test_pdf_docx_and_xlsx_are_read_with_their_own_parsers(api):
    import docx
    import fitz
    import openpyxl

    pdf = fitz.open()
    page = pdf.new_page()
    # 내장 한국어 글꼴 — 맑은 고딕 파일을 통째로 넣으면 10MB 제한을 넘는다(제한은 그대로 둔다)
    page.insert_text((72, 72), "1일차 2026-10-15\n09:00 경복궁\n12:30 토속촌삼계탕", fontname="korea", fontsize=12)
    pdf_bytes = pdf.tobytes()
    document = docx.Document()
    document.add_paragraph("1일차 2026-10-15")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "09:00", "경복궁"
    buffer = io.BytesIO()
    document.save(buffer)
    book = openpyxl.Workbook()
    book.active.append(["1일차", "2026-10-15"])
    book.active.append(["09:00", "경복궁"])
    sheet = io.BytesIO()
    book.save(sheet)

    client = _client()
    view = _send(client, _key(client), files=[("a.pdf", pdf_bytes), ("b.docx", buffer.getvalue()),
                                               ("c.xlsx", sheet.getvalue())])
    assert view["status"] == "review", view
    assert [s["kind"] for s in view["sources"]] == ["pdf", "docx", "xlsx"]
    assert _items(view, 0) == [("09:00", "경복궁"), ("12:30", "토속촌삼계탕")]
    assert _items(view, 1) == [("09:00", "경복궁")]
    assert _items(view, 2) == [("09:00", "경복궁")]


def test_unknown_files_are_refused_at_the_door_and_photos_need_a_reader(api):
    client = _client()
    headers = _key(client)
    refused = client.post("/v1/web/trip-intakes", headers=headers,
                          files=[("files", ("x.bin", b"\x00\x01\x02binary", "application/octet-stream"))])
    assert refused.status_code == 422 and refused.json()["error"]["code"] == "unsupported_format"
    empty = client.post("/v1/web/trip-intakes", headers=headers, data={"text": "  "})
    assert empty.status_code == 422 and empty.json()["error"]["code"] == "empty_intake"
    # 받아쓰기 모델이 없으면 사진은 「읽지 못함」— 지어내지 않는다
    view = _send(client, headers, files=[("plan.png", PHOTO.read_bytes())])
    assert view["status"] == "fatal" and view["fatal"]["code"] == "unsupported_format", view


def test_confirming_an_intake_that_could_not_be_read_is_a_409_not_a_server_error(api):
    """★「읽지 못했어요」 접수에 「등록하고 관리 시작」을 눌러도 서버 오류(500)가 나지 않는다 — 409 와 현재 상태를 돌려준다.
    오류 함수가 상세의 `status` 를 위치 인자 `status` 와 겹쳐 받아 `TypeError` 가 났다(그래서 위치 전용으로 바꿨다)."""
    client = _client()                                        # 받아쓰기 모델이 없어 사진은 읽지 못함(fatal)
    headers = _key(client)
    view = _send(client, headers, files=[("plan.png", PHOTO.read_bytes())])
    assert view["status"] == "fatal", view
    body = _confirm(client, headers, view["intake_id"], view["revision"], status=409)
    assert body["error"]["code"] == "intake_not_ready" and body["error"]["status"] == "fatal", body


def test_an_intake_is_yours_only(api):
    client = _client()
    mine = _key(client)
    other = _key(client)
    response = client.post("/v1/web/trip-intakes", headers=mine, data={"text": "09:00 경복궁"})
    intake = response.json()["intake_id"]
    assert client.get(f"/v1/web/trip-intakes/{intake}", headers=other).status_code == 404
    assert client.get(f"/v1/web/trip-intakes/{intake}").status_code == 401
    assert client.post("/v1/web/trip-intakes", data={"text": "09:00 경복궁"}).status_code == 401


@pytest.mark.parametrize("name", ["plan.PNG.pdf", "plan.txt"])
def test_the_kind_comes_from_the_bytes_not_the_name(api, name):
    """★확장자가 PDF 라도 첫 바이트가 PNG 면 사진이다."""
    vision = Vision()
    client = _client(vision)
    view = _send(client, _key(client), files=[(name, PHOTO.read_bytes())])
    assert view["sources"][0]["kind"] == "image" and _items(view) == WANT


# ── 2주차: 남은 줄(모델이 가리키기만) → 날짜 → 장소 ───────────────────────────
CHAT_PLAN = "1일차 2026-10-15\n09:00 경복궁\n점심은 토속촌에서 먹을래\n저녁엔 광장시장 가 보고 싶어"


class Labeler:
    """모델 흉내 — 남은 줄에 위치만 가리킨다. 하나는 일부러 원문을 바꿔 적는다(「토속툰」 — 실측에서 본 모양)."""

    def see(self, prompt, image):          # 사진은 안 쓴다
        raise AssertionError("사진이 없다")

    def json(self, system, user):
        assert "3: 점심은 토속촌에서 먹을래" in user and "2: 09:00" not in user   # ★규칙이 읽은 줄은 안 보낸다
        return {"spans": [{"line": 3, "quote": "점심", "role": "meal"},
                          {"line": 3, "quote": "토속촌", "role": "place"},
                          {"line": 4, "quote": "저녁", "role": "time"},
                          {"line": 4, "quote": "광장시장", "role": "place"},
                          {"line": 4, "quote": "광장시쟝", "role": "place"}]}


class Tour:
    PLACES = {"경복궁": ("126508", "12", 37.5796, 126.9770), "광장시장": ("264570", "38", 37.5700, 126.9996),
              "토속촌삼계탕": ("2717339", "39", 37.5778, 126.9716)}

    def __init__(self):
        self.asked = []

    def find(self, name, area_code=None):
        self.asked.append((name, area_code))
        if name not in self.PLACES:
            return None
        cid, ctype, lat, lon = self.PLACES[name]
        return {"matched_title": name, "content_id": cid, "content_type_id": ctype, "latitude": lat,
                "longitude": lon, "address": "서울특별시 종로구"}


class Kakao:
    def __init__(self):
        self.asked = []

    def search(self, query, size=5, near=None, **kw):     # ★운영 쪽 래퍼(`_KakaoNearHint`)가 `near` 를 넘긴다
        self.asked.append(query)
        return [{"id": "1", "name": "토속촌삼계탕", "category": "음식점 > 한식", "category_group": "FD6",
                 "address": "서울 종로구 자하문로5길 5", "latitude": 37.5778, "longitude": 126.9716}] \
            if query == "토속촌" else []


def test_unread_lines_are_pointed_at_by_the_model_and_places_are_looked_up(api):
    tour, kakao = Tour(), Kakao()
    client = TestClient(create_app(
        classifier=lambda _m: {"intent": "other", "issue_code": "other", "sentiment": "neutral"},
        domain_routers=[build_trip_router(chat_factory=Labeler, place_factory=lambda: tour,
                                          kakao_factory=lambda: kakao)]))
    view = _send(client, _key(client), text=CHAT_PLAN)
    items = view["sources"][0]["items"]
    titles = [i["fields"]["title"]["value"] for i in items]
    assert titles == ["경복궁", "토속촌", "광장시장"]
    # ★모델이 바꿔 적은 인용(「광장시쟝」)은 원문에 없어 버렸다 — 그 보고가 남는다
    report = view["sources"][0]["reading"]
    assert [r["quote"] for r in report["rejected"]] == ["광장시쟝"] and report["accepted"] == 4
    assert items[1]["fields"]["title"]["method"] == "llm_span"
    assert items[1]["fields"]["kind"]["value"] == "dining"                   # 「점심」 → 끼니
    # 날짜 — 본문에 적힌 날짜가 모델이 가리킨 줄까지 이어진다
    assert [i["date"] for i in items] == ["2026-10-15"] * 3
    # 장소 — 활동 CSV(관광공사 값) 정확 일치 · 활동이 아닌 곳(식당 「토속촌삼계탕」)은 카카오로 이름을 찾은 근거 그대로 둔다
    #   (`[2026-09-30]` 접수의 장소 확인이 활동 CSV 전용으로 바뀌어 식당을 관광공사로 다시 확인하지 않는다)
    places = [i["fields"]["place"] for i in items]
    assert [(p["value"]["name"], p["evidence"]["source"]) for p in places] == [
        ("경복궁", "tour_api"), ("토속촌삼계탕", "kakao"), ("광장시장", "tour_api")]
    assert places[1]["needs_review"] and not places[0]["needs_review"]      # 이름이 원문과 다르다 → 확인
    assert places[1]["value"]["kind"] == "dining"
    assert all(area == "1" for _, area in tour.asked)                        # ★서울 밖으로 새지 않는다
    assert kakao.asked == ["토속촌"]                                           # 카카오는 관광공사에 없을 때만


def test_a_model_outage_only_skips_the_pointing_step(api):
    class Down(Labeler):
        def json(self, system, user):
            raise TimeoutError("ollama down")

    client = TestClient(create_app(
        classifier=lambda _m: {"intent": "other", "issue_code": "other", "sentiment": "neutral"},
        domain_routers=[build_trip_router(chat_factory=Down)]))
    view = _send(client, _key(client), text=CHAT_PLAN)
    assert view["status"] == "review"
    items = view["sources"][0]["items"]
    assert [i["fields"]["title"]["value"] for i in items] == ["경복궁"]    # 규칙이 읽은 것은 그대로
    assert [line["read"] for line in view["sources"][0]["lines"]] == [True, True, False, False]
    assert view["sources"][0]["reading"]["error"].startswith("TimeoutError")   # ★장애는 이름으로 남는다


# ── 3주차: 확인 화면 → 고치기 → 등록 ─────────────────────────────────────
def _full_client():
    tour, kakao = Tour(), Kakao()
    client = TestClient(create_app(
        classifier=lambda _m: {"intent": "other", "issue_code": "other", "sentiment": "neutral"},
        domain_routers=[build_trip_router(chat_factory=Labeler, place_factory=lambda: tour,
                                          kakao_factory=lambda: kakao)]))
    return client, tour, kakao


def _stored(api, trip):
    """등록된 항목(내부 모양 — `detail` 포함). 공개 조회는 `detail` 을 싣지 않는다."""
    from app.modules.travel_ops.itinerary import TripStore

    with get_connection() as conn:
        return TripStore(api["tenant"]).latest(conn, UUID(trip["trip_id"]))[1]


def _edit(client, headers, view, *edits, status=200):
    response = client.post(f"/v1/web/trip-intakes/{view['intake_id']}/edits", headers=headers,
                           json={"revision": view["revision"], "edits": list(edits)})
    assert response.status_code == status, response.text
    return response.json()


def _confirm(client, headers, intake_id, revision, status=200):
    response = client.post(f"/v1/web/trip-intakes/{intake_id}/confirm", headers=headers, json={"revision": revision})
    assert response.status_code == status, response.text
    return response.json()


def test_the_confirm_screen_fixes_a_time_and_registers_one_trip(api):
    client, _, _ = _full_client()
    headers = _key(client)
    view = _send(client, headers, text=CHAT_PLAN)
    source = view["sources"][0]["source_id"]
    # 시각 없는 「점심」은 12:00, 「저녁엔 광장시장」은 18:00 — 원문의 끼니 말로 규칙 배치라고 적힌다
    assert view["check"]["ready"] is True, view["check"]["problems"]
    filled = {f["field"]: f["value"] for f in view["check"]["filled"] if f["field"].endswith("starts_at")}
    # (모델이 가리킨 항목은 규칙이 읽은 항목 뒤에 번호가 붙는다 — 경복궁 0 · 토속촌 1 · 광장시장 2)
    assert filled == {"items[1].starts_at": "12:00", "items[2].starts_at": "18:00"}
    # 고객이 광장시장을 저녁으로 고친다 → 새 판
    fixed = _edit(client, headers, view, {"source_id": source, "field": "items[2].starts_at", "value": "18:30"})
    assert fixed["revision"] == 2
    assert [f["field"] for f in fixed["check"]["filled"] if f["field"].endswith("starts_at")] == ["items[1].starts_at"]
    # ★낡은 판으로 고치거나 등록하면 409 — 다른 탭의 고침을 조용히 덮지 않는다
    _edit(client, headers, view, {"source_id": source, "field": "items[2].title", "value": "광장시장"}, status=409)
    stale = _confirm(client, headers, view["intake_id"], 1, status=409)
    assert stale["error"]["code"] == "stale_revision"
    done = _confirm(client, headers, view["intake_id"], 2)
    trip = done["trip"]
    assert done["status"] == "confirmed" and trip["created"] is True
    got = [(i["title"], i["starts_at"][11:16], i["kind"]) for i in trip["items"]]
    assert got == [("경복궁", "09:00", "activity"), ("토속촌", "12:00", "dining"), ("광장시장", "18:30", "activity")]
    # 근거가 등록까지 따라간다 — 규칙으로 채운 시각 · 모델이 가리킨 원문 · 조회 출처
    lunch = _stored(api, trip)[1].detail["provenance"]
    assert lunch["starts_at"]["method"] == "rule" and lunch["title"]["method"] == "llm_span"
    assert lunch["place"]["method"] == "lookup"
    # ★관광공사에서 온 장소는 그 여행 전용 행이다(공용 표에 쌓지 않는다)
    with get_connection() as conn, conn.cursor() as cur:
        cur.execute("SELECT count(*) FILTER (WHERE trip_scope IS NULL), count(*) FILTER (WHERE trip_scope::text=%s) "
                    "FROM places WHERE tenant_id=%s", (trip["trip_id"], api["tenant"]))
        assert cur.fetchone() == (0, 3)
    # 같은 확인을 한 번 더 눌러도 여행은 하나다
    again = _confirm(client, headers, view["intake_id"], 2)
    assert again["trip"]["trip_id"] == trip["trip_id"] and again["trip"]["created"] is False
    after = client.get(f"/v1/web/trip-intakes/{view['intake_id']}", headers=headers).json()
    assert after["status"] == "confirmed" and after["trip_id"] == trip["trip_id"]


def test_a_plan_without_dates_asks_for_the_first_day_before_registering(api):
    client, _, _ = _full_client()
    headers = _key(client)
    view = _send(client, headers, text="1일차\n09:00 경복궁\n2일차\n10:00 광장시장")
    blocked = _confirm(client, headers, view["intake_id"], view["revision"], status=422)
    assert blocked["error"]["code"] == "intake_incomplete"
    assert {p["code"] for p in blocked["error"]["problems"]} == {"no_date"}
    fixed = _edit(client, headers, view, {"field": "trip.first_day", "value": "2026-11-02"})
    assert fixed["check"]["ready"] is True
    trip = _confirm(client, headers, view["intake_id"], fixed["revision"])["trip"]
    assert [i["starts_at"][:16] for i in trip["items"]] == ["2026-11-02T09:00", "2026-11-03T10:00"]


def test_an_unfound_place_is_fixed_by_name_or_left_without_a_place(api):
    client, _, _ = _full_client()
    headers = _key(client)
    view = _send(client, headers, text="1일차 2026-10-15\n09:00 경복궁\n11:00 동네 산책\n13:00 모르는식당")
    source = view["sources"][0]["source_id"]
    problems = {(p["code"], p["field"]) for p in view["check"]["problems"]}
    assert problems == {("no_place", "items[1].place"), ("no_place", "items[2].place")}
    # 이름으로 고치면 다시 찾는다. 못 찾는 이름은 받지 않는다(지어내지 않는다)
    refused = _edit(client, headers, view, {"source_id": source, "field": "items[2].place",
                                            "value": {"name": "없는가게"}}, status=422)
    assert refused["error"]["code"] == "place_not_found"
    fixed = _edit(client, headers, view,
                  {"source_id": source, "field": "items[2].place", "value": {"name": "광장시장"}},
                  {"source_id": source, "field": "items[1].place", "value": {"none": True}})
    assert fixed["check"]["ready"] is True
    trip = _confirm(client, headers, view["intake_id"], fixed["revision"])["trip"]
    walk = next(i for i in trip["items"] if i["title"] == "동네 산책")
    assert walk["place"] is None


def test_a_booking_number_is_carried_and_protects_the_item(api):
    from app.modules.travel_ops.pending import protected_reason

    client, _, _ = _full_client()
    headers = _key(client)
    view = _send(client, headers, text="1일차 2026-10-15\n19:00 경복궁 · 예약번호 KY-20931")
    trip = _confirm(client, headers, view["intake_id"], view["revision"])["trip"]
    items = _stored(api, trip)
    assert items[0].detail["booking"] == {"booking_no": "KY-20931", "declared_by": "customer_plan"}
    assert protected_reason(items[0]) == "booked"          # ★바꾸기 전에 묻는다


LUNCH_PLAN = "\n".join(["1일차 2026-10-15", "09:00 경복궁", "점심은 토속촌에서 먹을래", "15:00 광장시장"])


def test_a_lunch_line_is_placed_at_lunch_even_after_a_later_timed_stop(api):
    """☆2026-09-27 실제 화면: 「점심은 토속촌에서」가 18:00 에 놓였다 — 모델이 가리킨 항목은 번호가 규칙 항목 뒤에
    붙어서, 번호 순서로 채우면 오후 일정 뒤로 밀린다. 원문 줄 순서 + 원문의 끼니 말로 채운다."""
    client, _, _ = _full_client()
    headers = _key(client)
    view = _send(client, headers, text=LUNCH_PLAN)
    filled = [f for f in view["check"]["filled"] if f["field"].endswith("starts_at")]
    assert [(f["field"], f["value"]) for f in filled] == [("items[2].starts_at", "12:00")]
    assert "「점심」" in filled[0]["note"]


ALIAS_FIRST = chr(10).join(["1일차 2026-10-15", "09:00 경복궁", "11:00 옛날시장"])
ALIAS_NEXT = chr(10).join(["1일차 2026-10-16", "10:00 옛날시장"])


def test_a_place_name_the_customer_fixed_is_remembered_for_the_next_plan(api):
    """★고객이 고친 (원래 글 → 고친 글)을 별칭으로 쌓는다(030) — 다음 계획의 같은 말은 그 이름으로 다시 찾는다.
    둘 다 고객 글이다 — 관광공사·카카오 값은 별칭 표에 넣지 않는다."""
    client, _, _ = _full_client()
    headers = _key(client)
    first = _send(client, headers, text=ALIAS_FIRST)
    source = first["sources"][0]["source_id"]
    assert ("no_place", "items[1].place") in {(p["code"], p["field"]) for p in first["check"]["problems"]}
    _edit(client, headers, first, {"source_id": source, "field": "items[1].place", "value": {"name": "광장시장"}})
    with get_connection() as conn, conn.cursor() as cur:
        cur.execute("SELECT phrase, replacement, source FROM place_aliases WHERE tenant_id=%s", (api["tenant"],))
        assert cur.fetchall() == [("옛날시장", "광장시장", "customer")]
    again = _send(client, headers, text=ALIAS_NEXT)
    place = again["sources"][0]["items"][0]["fields"]["place"]
    assert place["value"]["name"] == "광장시장" and place["needs_review"] is True
    assert place["evidence"]["tried"][0] == "alias:옛날시장→광장시장"


def test_the_survey_sent_with_confirm_is_checked_and_stays_on_the_trip(api):
    """★`[2026-09-28]` 「등록하고 관리 시작」에도 설문을 싣는다 — `/v1/web/trips` 와 같은 검사(`_create_trip`)."""
    from app.modules.travel_ops.survey import SURVEY_VERSION

    client, _, _ = _full_client()
    headers = _key(client)
    view = _send(client, headers, text=CHAT_PLAN)
    url = f"/v1/web/trip-intakes/{view['intake_id']}/confirm"
    bad = client.post(url, headers=headers, json={"revision": view["revision"],
                                                  "survey": {"version": SURVEY_VERSION, "unknown": 1}})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "invalid_survey", bad.text
    done = client.post(url, headers=headers, json={"revision": view["revision"],
                                                   "survey": {"version": SURVEY_VERSION, "on_disruption": "ask_first"}})
    assert done.status_code == 200, done.text
    with get_connection() as conn, conn.cursor() as cur:
        cur.execute("SELECT constraints FROM trips WHERE tenant_id=%s AND trip_id=%s",
                    (api["tenant"], done.json()["trip"]["trip_id"]))
        assert cur.fetchone()[0]["survey"]["on_disruption"] == "ask_first"

